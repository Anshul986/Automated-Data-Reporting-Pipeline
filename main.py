"""
etl_pipeline.py

Extract -> Transform -> Load pipeline for weekly hunt reports.

- Extract: read the raw .xlsx export into a DataFrame.
- Transform: drop "Total" rows, filter out players who haven't been in the
  guild long enough (or have no hunt-date data), clean/normalize the
  Points and Goal Percentage columns, and sort by performance.
- Load: write the cleaned data to a formatted, color-coded .xlsx report.
"""

import os
import datetime
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.utils import get_column_letter

REQUIRED_COLUMNS = ["Name", "Points (Hunt)", "Goal Percentage (Hunt)"]


# ---------------------------------------------------------------------------
# TRANSFORM: guild-presence filtering
# ---------------------------------------------------------------------------
def filter_players_by_guild_presence(df):
    filtered_players = []

    try:
        filtered_df = df.copy()

        # Remove "Total" row from Name column first
        original_count_before_total = len(filtered_df)
        filtered_df = filtered_df[filtered_df['Name'].str.lower() != 'total']
        total_removed = original_count_before_total - len(filtered_df)
        if total_removed > 0:
            print(f"Removed {total_removed} 'Total' row(s)")

        # Check if we have the required date columns
        date_cols_present = all(col in filtered_df.columns for col in ["First Hunt Time", "Last Hunt Time"])

        if not date_cols_present:
            print("Date columns not found - skipping date-based filtering")
            return filtered_df, []

        # Convert date columns to datetime
        filtered_df['First Hunt Time'] = pd.to_datetime(filtered_df['First Hunt Time'], errors='coerce')
        filtered_df['Last Hunt Time'] = pd.to_datetime(filtered_df['Last Hunt Time'], errors='coerce')

        # Identify players with valid dates for filtering
        valid_dates_mask = filtered_df['First Hunt Time'].notna() & filtered_df['Last Hunt Time'].notna()
        players_with_valid_dates = filtered_df[valid_dates_mask].copy()
        players_without_dates = filtered_df[~valid_dates_mask].copy()

        print(f"Players with valid dates: {len(players_with_valid_dates)}")
        print(f"Players without valid dates (0% hunters): {len(players_without_dates)}")

        # Calculate days in guild only for players with valid dates
        players_with_valid_dates['Days In Guild'] = (players_with_valid_dates['Last Hunt Time'] - players_with_valid_dates['First Hunt Time']).dt.days.abs()

        # Filter out players with less than 4 days in guild (only from those with valid dates)
        players_to_filter = players_with_valid_dates[players_with_valid_dates['Days In Guild'] < 4]

        # Also include players without dates (0% hunters) in the filtered list
        players_without_dates['Days In Guild'] = 0  # Mark as 0 days for tracking
        players_without_dates['Reason'] = '0% hunter (no date data)'

        # Combine both filtered groups
        all_filtered_players = pd.concat([players_to_filter, players_without_dates], ignore_index=True)

        # Prepare filtered players info for display
        filtered_players_info = []
        for _, player in all_filtered_players.iterrows():
            player_info = {
                'Name': player['Name'],
                'Days In Guild': player.get('Days In Guild', 0),
                'Reason': player.get('Reason', 'Less than 4 days in guild')
            }
            filtered_players_info.append(player_info)

        # Keep ONLY players with 4+ days in guild (exclude both <4 days AND 0% hunters)
        players_with_enough_days = players_with_valid_dates[players_with_valid_dates['Days In Guild'] >= 4]

        # Final dataframe contains ONLY players with 4+ days in guild
        final_df = players_with_enough_days.copy()

        # Drop the temporary column
        if 'Days In Guild' in final_df.columns:
            final_df = final_df.drop('Days In Guild', axis=1)

        print(f"Original count: {original_count_before_total}")
        print(f"After removing 'Total': {len(filtered_df)}")
        print(f"Filtered out {len(players_to_filter)} players with less than 4 days in guild")
        print(f"Filtered out {len(players_without_dates)} 0% hunters (no date data)")
        print(f"Final count (only players with 4+ days in guild): {len(final_df)}")

        return final_df, filtered_players_info

    except Exception as e:
        print(f"Error in date filtering: {e}")
        return df, []


# ---------------------------------------------------------------------------
# TRANSFORM: value parsing / cleaning
# ---------------------------------------------------------------------------
def parse_percent_cell(raw):
    if raw is None: return None
    if isinstance(raw, (int, float)): return float(raw)
    s = str(raw).strip().replace(",", "").replace(" ", "")
    if not s or s.lower() in ("-", "--", "n/a", "na"): return None
    if s.endswith("%"): s = s[:-1]
    try: return float(s)
    except: return None


def convert_and_round_series(series):
    parsed = series.map(parse_percent_cell)
    numeric = parsed.dropna().astype(float)
    if numeric.empty:
        return pd.Series(0, index=series.index)

    if numeric.max() <= 100:
        numeric *= 100

    rounded = numeric.round().astype(int)

    def get_value(i):
        if not pd.isna(parsed[i]):
            return int(rounded.get(i, 0))
        else:
            return 0

    return parsed.index.to_series().map(get_value)


# ---------------------------------------------------------------------------
# LOAD: Excel formatting / color coding
# ---------------------------------------------------------------------------
def format_and_color_excel(path):
    wb = load_workbook(path)
    ws = wb.active
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")

    widths = {1: 30, 2: 18, 3: 18, 4: 40}
    for i in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(i, 18)

    # Color coding for top 3 ranks
    rank_colors = {
        1: "e79d5e",  # First rank color
        2: "5e95e7",  # Second rank color
        3: "faa3d1"   # Third rank color
    }

    # Find Kidnape me row
    kidnape_row = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value and "kidnape me" in str(ws.cell(r, 1).value).lower():
            kidnape_row = r
            break

    # Add rank numbers to issues column for top 3
    for r in range(2, min(5, ws.max_row + 1)):
        rank = r - 1
        issue_cell = ws.cell(r, 4)
        if rank == 1:
            issue_cell.value = "Rank 1"
        elif rank == 2:
            issue_cell.value = "Rank 2"
        elif rank == 3:
            issue_cell.value = "Rank 3"
        # Keep the text but no background color
        issue_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Apply colors to top 3 rows
    for r in range(2, min(5, ws.max_row + 1)):
        rank = r - 1
        if rank in rank_colors:
            # Set the color for the row (only first 3 columns)
            fill = PatternFill(start_color=rank_colors[rank], end_color=rank_colors[rank], fill_type="solid")

            # Apply color to first 3 columns only
            for c in range(1, 4):
                cell = ws.cell(r, c)
                cell.fill = fill
                cell.alignment = Alignment(horizontal="left" if c == 1 else "center", vertical="center")

            # Issues column - no color, left aligned (already set above)

    # Handle Kidnape me row
    if kidnape_row:
        # Make Kidnape me row white (first 3 columns)
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        for c in range(1, 4):
            kidnape_cell = ws.cell(kidnape_row, c)
            kidnape_cell.fill = white_fill
            kidnape_cell.alignment = Alignment(horizontal="left" if c == 1 else "center", vertical="center")

        # Add BANK text to issues column for Kidnape me row (no color, left aligned)
        bank_issue_cell = ws.cell(kidnape_row, 4)
        bank_issue_cell.value = "BANK"
        bank_issue_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Format the rest of the rows (after top 3)
    start_row = 5
    if kidnape_row and kidnape_row < 5:  # Adjust if Kidnape me is in top rows
        start_row = 6

    for r in range(start_row, ws.max_row + 1):
        # Skip Kidnape me row (already formatted)
        if kidnape_row and r == kidnape_row:
            continue

        try:
            # Extract numeric value from percentage string (e.g., "150%" -> 150)
            cell_value = ws.cell(r, 3).value
            if cell_value and isinstance(cell_value, str) and cell_value.endswith('%'):
                p = float(cell_value.rstrip('%'))
            else:
                p = float(cell_value or 0)
        except:
            p = 0

        if p >= 100: color = "90EE90"
        elif p >= 90: color = "F5DEB3"
        elif p >= 70: color = "FFFF99"
        else: color = "FFB3B3"
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

        for c in range(1, 4):
            cell = ws.cell(r, c)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="left" if c == 1 else "center", vertical="center")

        # Issues column - no color, left aligned
        issue_cell = ws.cell(r, 4)
        issue_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Remove "Issues" header and replace with empty string
    ws.cell(1, 4).value = ""

    wb.save(path)


# ---------------------------------------------------------------------------
# PIPELINE: extract / transform / load orchestration
# ---------------------------------------------------------------------------
def extract(input_path: str) -> pd.DataFrame:
    """Read the raw .xlsx export into a DataFrame."""
    return pd.read_excel(input_path)


def transform(df: pd.DataFrame):
    """
    Clean and filter the raw hunt data.

    Returns:
        (clean_df, stats) where stats holds counts/details of what was
        removed during filtering.
    """
    if not all(c in df.columns for c in REQUIRED_COLUMNS):
        missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
        raise ValueError(f"Missing required columns: {missing}")

    total_removed_count = len(df[df['Name'].str.lower() == 'total'])

    filtered_df, filtered_players_info = filter_players_by_guild_presence(df)

    clean_df = filtered_df[REQUIRED_COLUMNS].copy()
    clean_df["Issues"] = ""
    clean_df["Points (Hunt)"] = pd.to_numeric(clean_df["Points (Hunt)"], errors="coerce").fillna(0).astype(int)

    # Convert goal percentage and add % sign
    goal_percentages = convert_and_round_series(clean_df["Goal Percentage (Hunt)"])
    clean_df["Goal_Percentage_Numeric"] = goal_percentages.clip(lower=0)
    clean_df["Goal Percentage (Hunt)"] = clean_df["Goal_Percentage_Numeric"].astype(str) + "%"

    # Sort by the numeric column, then drop the helper column
    clean_df = clean_df.sort_values("Goal_Percentage_Numeric", ascending=False).reset_index(drop=True)
    clean_df = clean_df.drop("Goal_Percentage_Numeric", axis=1)

    stats = {
        "total_rows_removed": total_removed_count,
        "filtered_players": filtered_players_info,
    }
    return clean_df, stats


def load(df: pd.DataFrame, output_path: str):
    """Write the cleaned DataFrame to a formatted, color-coded .xlsx report."""
    df.to_excel(output_path, index=False)
    format_and_color_excel(output_path)


def run_pipeline(input_path: str, output_dir: str = "downloads") -> str:
    """
    Run the full Extract -> Transform -> Load pipeline on a raw hunt export.

    Returns the path to the generated report.
    """
    os.makedirs(output_dir, exist_ok=True)

    df = extract(input_path)
    clean_df, stats = transform(df)

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = os.path.join(output_dir, f"Processed_Report_{ts}.xlsx")
    load(clean_df, output_path)

    print(f"Removed {stats['total_rows_removed']} 'Total' row(s)")
    print(f"Filtered out {len(stats['filtered_players'])} player(s)")
    print(f"Report written to {output_path}")

    return output_path


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("Usage: python etl_pipeline.py <input_file.xlsx> [output_dir]")
        sys.exit(1)

    input_file = sys.argv[1]
    out_dir = sys.argv[2] if len(sys.argv) > 2 else "downloads"
    run_pipeline(input_file, out_dir)
